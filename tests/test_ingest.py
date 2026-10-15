from io import BytesIO
from zipfile import ZipFile

import pymupdf
import pytest
from docx import Document
from fastapi import UploadFile
from starlette.datastructures import Headers

from app.services.chunking import chunk_sections
from app.services.errors import ServiceError
from app.services.ingest import TextSection, extract_text, save_upload, stored_path


def pdf_bytes(text="DEMONSTRACAO FICTICIA. Consulte a disponibilidade para early check-in."):
    with pymupdf.open() as pdf:
        pdf.new_page()
        pdf.new_page()
        pdf.new_page().insert_text((40, 50), text)
        return pdf.tobytes()


def test_pdf_extraction_preserves_page(tmp_path):
    path = tmp_path / "procedure.pdf"
    path.write_bytes(pdf_bytes())
    sections = extract_text(path)
    assert len(sections) == 1
    assert sections[0].page == 3
    assert "Consulte a disponibilidade" in sections[0].text


def test_docx_preserves_sections_and_tables(tmp_path):
    doc = Document()
    doc.add_heading("Early Check-in", level=1)
    doc.add_paragraph("Consultar a disponibilidade.")
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Procedimento"
    table.cell(0, 1).text = "Consultar recepção"
    doc.add_heading("Exceções", level=1)
    doc.add_paragraph("Não confirmar sem consultar.")
    path = tmp_path / "procedure.docx"
    doc.save(path)
    sections = extract_text(path)
    assert [item.section for item in sections] == ["Early Check-in", "Exceções"]
    assert all(item.page == 0 for item in sections)
    assert "Procedimento | Consultar recepção" in sections[0].text


def test_chunking_preserves_paragraphs_and_overlap():
    paragraphs = [f"Parágrafo {i}: " + "palavra " * 80 for i in range(20)]
    chunks = chunk_sections([TextSection("\n\n".join(paragraphs), page=3)], 1000, 150)
    assert len(chunks) > 1
    assert all(chunk.page == 3 and len(chunk.text) <= 4000 for chunk in chunks)
    assert all(any(paragraph.strip() in chunk.text for chunk in chunks) for paragraph in paragraphs)
    assert chunks[0].text[-100:] in chunks[1].text


def test_chunking_long_paragraph_without_spaces_and_page_isolation():
    chunks = chunk_sections([TextSection("x" * 9000, page=1), TextSection("Segunda página", page=2)])
    assert max(len(chunk.text) for chunk in chunks) <= 4000
    assert chunks[-1].page == 2
    assert chunks[-1].text == "Segunda página"
    assert chunk_sections([TextSection("  ")]) == []
    with pytest.raises(ValueError):
        chunk_sections([], 10, 10)


@pytest.mark.parametrize("name, mime, data, status", [
    ("../../file.pdf", "application/pdf", b"%PDF-", 422),
    ("file.exe", "application/octet-stream", b"bad", 415),
    ("file.pdf", "text/html", b"bad", 415),
    ("file.pdf", "application/pdf", b"not a PDF", 422),
    ("file.docx", "application/octet-stream", b"not a ZIP", 422),
    ("file.pdf", "application/pdf", b"", 422),
    ("file.pdf", "application/pdf", b"%PDF-" + b"x" * 100, 413),
])
def test_invalid_upload_is_rejected_without_residue(tmp_path, name, mime, data, status):
    upload = UploadFile(BytesIO(data), filename=name, headers=Headers({"content-type": mime}))
    with pytest.raises(ServiceError) as error:
        save_upload(upload, tmp_path, 100)
    assert error.value.status_code == status
    assert not list(tmp_path.iterdir())


def test_upload_name_and_path_are_safe(tmp_path):
    upload = UploadFile(BytesIO(pdf_bytes()), filename="POP_EARLY_CHECKIN.pdf", headers=Headers({"content-type": "application/pdf"}))
    path, name = save_upload(upload, tmp_path, 100000)
    assert name == "POP_EARLY_CHECKIN.pdf"
    assert path.name != name
    assert stored_path(tmp_path, path.name) == path.resolve()
    with pytest.raises(ServiceError):
        stored_path(tmp_path, "../outside.pdf")


def test_empty_pdf_and_invalid_docx_container(tmp_path):
    path = tmp_path / "empty.pdf"
    with pymupdf.open() as doc:
        doc.new_page()
        doc.save(path)
    with pytest.raises(ServiceError, match="OCR"):
        extract_text(path)
    path = tmp_path / "invalid.docx"
    with ZipFile(path, "w") as archive:
        archive.writestr("other.xml", "invalid")
    with pytest.raises(ServiceError):
        extract_text(path)
