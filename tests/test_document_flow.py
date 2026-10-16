import json
from contextlib import contextmanager
from dataclasses import replace
from io import BytesIO

import httpx
import pymupdf
import pytest
from django.core.files.uploadedfile import SimpleUploadedFile
from docx import Document

from app import runtime
from app.models import Documento, Setor, Treinamento
from app.services.documents import DocumentService
from app.services.embeddings import EmbeddingService
from app.services.llm import LLMService
from app.services.rag import NOT_FOUND, RagService
from app.services.vector_store import VectorStore

pytestmark = pytest.mark.django_db


def make_pdf():
    with pymupdf.open() as pdf:
        pdf.new_page()
        pdf.new_page()
        pdf.new_page().insert_text((40,50),"DEMONSTRACAO FICTICIA. Early check-in: consultar disponibilidade antes de confirmar.")
        return pdf.tobytes()


@pytest.fixture
def flow(staff_client,platform_settings,settings,monkeypatch):
    config = replace(platform_settings,google_api_key="offline-google",openrouter_api_key="offline-router",embedding_dimensions=3)
    settings.SERVICE_SETTINGS = config
    store = VectorStore(config)
    calls = []
    state = {"fail_embeddings":False}

    def handler(request):
        data = json.loads(request.content)
        calls.append((str(request.url),data))
        if "googleapis.com" in str(request.url):
            if state["fail_embeddings"]:
                return httpx.Response(429,json={"error":"limit"})
            vectors = [{"values":[0,0,1] if "salário" in item["content"]["parts"][0]["text"].lower() else [1,0,0]} for item in data["requests"]]
            return httpx.Response(200,json={"embeddings":vectors})
        return httpx.Response(200,json={"choices":[{"message":{"content":"Consulte a disponibilidade antes de confirmar [1]."}}]})

    @contextmanager
    def services():
        with httpx.Client(transport=httpx.MockTransport(handler)) as client:
            embeddings = EmbeddingService(config,client)
            yield DocumentService(config,embeddings,store), RagService(embeddings,store,LLMService(config,client),config.max_distance)

    monkeypatch.setattr(runtime,"services",services)
    sector = Setor.objects.create(nome="Recepção")
    for title in ("Early Check-in","Financeiro"):
        Treinamento.objects.create(setor=sector,titulo=title)
    return staff_client,config,calls,state,store


def upload(client,training=1):
    file = SimpleUploadedFile("POP_EARLY_CHECKIN.pdf",make_pdf(),content_type="application/pdf")
    return client.post(f"/admin/treinamentos/{training}/documentos",{"arquivo":file})


def ask(client,question="Como proceder?"):
    return client.post("/treinamentos/1/perguntar",{"pergunta":question},content_type="application/json")


def test_pdf_upload_chat_scope_and_real_chroma(flow):
    client,config,calls,_,store = flow
    assert upload(client).status_code == 302
    assert upload(client,2).status_code == 302
    result = ask(client,"O que fazer quando o hóspede chega antes do horário?").json()
    assert "disponibilidade" in result["resposta"]
    assert len(result["fontes"]) == 1
    assert result["fontes"][0]["pagina"] == 3
    assert result["fontes"][0]["documento_id"] == 1
    llm_calls = [data for url,data in calls if "openrouter.ai" in url]
    assert len(json.loads(llm_calls[0]["messages"][1]["content"])["CONTEXTO"]) == 1
    response = client.get("/documentos/1/arquivo")
    assert response.status_code == 200
    assert b"".join(response.streaming_content).startswith(b"%PDF-")
    assert client.post("/admin/documentos/1/reindexar").status_code == 302
    assert store.collection.count() == 2
    assert "Indexado" in client.get("/admin/treinamentos/1/documentos").content.decode()


def test_missing_and_unrelated_documents_abstain(flow):
    client,_,calls,_,_ = flow
    assert ask(client).json() == {"resposta":NOT_FOUND,"fontes":[]}
    assert not calls
    upload(client)
    assert ask(client,"Qual é o salário?").json() == {"resposta":NOT_FOUND,"fontes":[]}
    assert not any("openrouter.ai" in url for url,_ in calls)


def test_failed_index_keeps_pending_file_and_retries(flow):
    client,config,calls,state,_ = flow
    state["fail_embeddings"] = True
    response = upload(client)
    assert response.status_code == 429
    assert response["Retry-After"] == "60"
    document = Documento.objects.get(pk=1)
    assert not document.processado
    assert (config.documents_path / document.caminho).is_file()
    before = len(calls)
    assert ask(client).json()["resposta"] == NOT_FOUND
    assert len(calls) == before
    state["fail_embeddings"] = False
    assert client.post("/admin/documentos/1/reindexar").status_code == 302
    document.refresh_from_db()
    assert document.processado
    state["fail_embeddings"] = True
    assert client.post("/admin/documentos/1/reindexar").status_code == 429
    assert ask(client).json()["resposta"] == NOT_FOUND


def test_delete_removes_metadata_file_and_vectors(flow):
    client,config,_,_,store = flow
    upload(client)
    path = config.documents_path / Documento.objects.get(pk=1).caminho
    assert client.post("/admin/documentos/1/excluir").status_code == 302
    assert not path.exists()
    assert not Documento.objects.exists()
    assert store.collection.count() == 0
    assert client.get("/documentos/1/arquivo").status_code == 404


def test_docx_tables_sections_and_chat(flow):
    client,_,_,_,_ = flow
    doc = Document()
    doc.add_heading("Early Check-in",level=1)
    doc.add_paragraph("Consultar disponibilidade antes de confirmar.")
    stream = BytesIO()
    doc.save(stream)
    file = SimpleUploadedFile("POP.docx",stream.getvalue(),content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document")
    assert client.post("/admin/treinamentos/1/documentos",{"arquivo":file}).status_code == 302
    result = ask(client).json()
    assert result["fontes"][0]["pagina"] is None
    assert result["fontes"][0]["secao"] == "Early Check-in"


def test_endpoint_guards(flow):
    client,_,_,_,_ = flow
    assert client.post("/treinamentos/999/perguntar",{"pergunta":"Teste"},content_type="application/json").status_code == 404
    assert client.post("/treinamentos/1/perguntar",data="bad",content_type="application/json").status_code == 422
    assert client.post("/treinamentos/1/perguntar",{"pergunta":"Teste"}).status_code == 415
    assert client.generic("POST","/admin/treinamentos/1/documentos",b"",content_type="multipart/form-data; boundary=test",CONTENT_LENGTH=str(20*1024*1024)).status_code == 413
